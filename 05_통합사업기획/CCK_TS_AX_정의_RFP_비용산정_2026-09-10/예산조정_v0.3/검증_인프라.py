from pathlib import Path
from zipfile import ZipFile
import json
from docx import Document

b=Path(__file__).resolve().parent
d=json.loads((b/'사업정의_데이터.json').read_text(encoding='utf-8'))
old=json.loads((b.parent/'예산조정_v0.2/사업정의_데이터.json').read_text(encoding='utf-8'))
assert d['infrastructure_investment']==0
assert [r[0] for r in d['direct']]==['D02','D03','D04','D05','D06','D07','D08','D09','D10','D11']
assert [r[0] for r in d['opsdirect']]==['OD01','OD03']
assert d['wps']==old['wps']
assert {r['id'] for r in d['requirements']}=={r['id'] for r in old['requirements']}
assert old['budget']['1단계']['gross']-d['budget']['1단계']['gross']==33000000
assert old['budget']['2단계']['gross']==d['budget']['2단계']['gross']
assert old['budget']['운영']['gross']-d['budget']['운영']['gross']==4400005
for p in b.glob('0[1-4]*.docx'):
    doc=Document(p)
    text='\n'.join([x.text for x in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert '0원' in text,p.name
    for bad in ['신규 구매 필요 시','견적 반영 후 범위 재조정','기존 GPU 제공 조건']:
        assert bad not in text,(p.name,bad)
for p in [b/'00_산출물_안내.html',b/'05_비용산정_설명.md']:
    text=p.read_text(encoding='utf-8')
    assert '신규 인프라 투자 0원' in text
    assert '1,118,387,104' in text
with ZipFile(b/'CCK_TS_AX_예산조정_패키지.zip') as z:
    assert len(z.namelist())==14
    for name in z.namelist():assert z.read(name)==(b/name).read_bytes(),name
report={'infrastructure_zero':True,'infrastructure_items_removed':True,'scope_and_effort_unchanged':True,'phase1_gross_reduction':33000000,'annual_ops_gross_reduction':4400005,'current_docx_count':4,'package_files_matched':14}
(b/'검증/인프라제약_검증.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
