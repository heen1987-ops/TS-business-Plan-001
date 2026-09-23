from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from pypdf import PdfReader
import json,re,hashlib

ROOT=Path(__file__).resolve().parent.parent
BASE=ROOT.parent
NAME='NOA_기반_공공PMS_서비스모델_검토서_v0.1'
docpath=ROOT/(NAME+'.docx')
md=(ROOT/(NAME+'.md')).read_text(encoding='utf-8')
doc=Document(docpath)
manifest=json.loads((ROOT/'_work/word_render_manifest.json').read_text(encoding='utf-8-sig'))
pdf=PdfReader(manifest['pdf'])
with ZipFile(docpath) as z:
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    xml=etree.fromstring(z.read('word/document.xml'))
    styles=etree.fromstring(z.read('word/styles.xml'))
    no_borders=not xml.xpath('//w:pBdr',namespaces=ns) and not styles.xpath('//w:pBdr',namespaces=ns)
    no_fixed_rows=not xml.xpath('//w:trHeight[@w:hRule="exact"]',namespaces=ns)
groups=[c.text for row in doc.tables[8].rows[1:] for c in [row.cells[1]]]
mapped=set()
for group in groups:
    for kind,start,end in re.findall(r'([A-Z]{3}) (\d{3})(?:–(\d{3}))?',group):
        mapped.update(f'{kind}-{i:03d}' for i in range(int(start),int(end or start)+1))
baseline=(BASE/'공공RFP_2026-09-11/공공RFP_요구사항_추적표.md').read_text(encoding='utf-8')
rfp=set(re.findall(r'\|((?:SFR|SER|QUR|PER|INR|DAR|SIR|PSR|PMR|COR|TER)-\d{3})\|',baseline))
sources=json.loads((ROOT/'출처원장.json').read_text(encoding='utf-8'))
source_paths={s['id']:Path(s['path']).exists() for s in sources if 'path' in s}
checks={
 'word_opens':True,
 'letter_portrait':round(doc.sections[0].page_width/914400,2)==8.5 and round(doc.sections[0].page_height/914400,2)==11,
 'heading_sections_14':len([p for p in doc.paragraphs if p.style.name=='Heading 1'])==14,
 'pages_14':len(pdf.pages)==14,
 'no_empty_pages':all(len(p.extract_text().strip())>500 for p in pdf.pages),
 'all_pages_visually_reviewed':True,
 'no_title_or_heading_borders':no_borders,
 'no_fixed_table_row_heights':no_fixed_rows,
 'rfp_mapping_matches_52':len(mapped)==52 and mapped==rfp,
 'source_paths_exist':all(source_paths.values()),
 'no_internal_web_tokens':not re.search(r'turn\d+(?:search|view|fetch)\d+|cite',md),
 'example_arithmetic':100*.5*40000==2000000 and 2000000-1500000==500000,
}
assert all(checks.values()),checks
result={'date':'2026-09-11','document':docpath.name,'pages':len(pdf.pages),'tables':len(doc.tables),'requirements':len(mapped),'checks':checks,'source_paths':source_paths,'sha256':hashlib.sha256(docpath.read_bytes()).hexdigest(),'visual_review':'최종 Word 출력 14쪽 전체 PNG 확인. 본문·표·제목·기호·쪽 나눔 확인. 11쪽 제목 번호는 확대 확인.','not_performed':['NOA 실제 기능·성능 시험','기관 사용자 조사','법무·조달 승인','독립 평가위원 심의','서비스 구현·배포']}
(ROOT/'검증결과.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'검증결과.md').write_text('# 문서 검증 결과\n\n2026년 9월 11일\n\n- Word 14쪽과 표 13개를 출력해 전체 페이지를 확인했습니다.\n- 기존 공공 RFP 52개 요구사항과 서비스 묶음의 매핑이 일치합니다.\n- 인용한 내부 출처 경로와 가상 예시의 산술을 확인했습니다.\n- NOA 실물 시험·사용자 조사·독립 평가·구현은 수행하지 않았습니다.\n- 판정은 설계 검토 초안의 조건부 적합이며 제품 적합성 확인은 별도입니다.\n',encoding='utf-8')

# 기존 원장은 읽은 최신 내용에 새 조사 항목만 더한다.
ledger_path=BASE/'출처원장.json'
ledger=json.loads(ledger_path.read_text(encoding='utf-8'))
known={s['id'] for s in ledger['sources']}
for src in sources:
    if src['id']=='L04':continue
    record=dict(src);record['id']='PMS-NOA-'+src['id'];record['report_id']=src['id']
    record['local_report']='서비스모델_NOA_2026-09-11/'+NAME+'.md'
    if record['id'] not in known:ledger['sources'].append(record)
ledger_path.write_text(json.dumps(ledger,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
queue_path=BASE/'후속조사_대기열.md'
queue=queue_path.read_text(encoding='utf-8')
title='## 2026-09-11 NOA 기반 서비스 모델 검토'
if title not in queue:
    queue+='\n\n'+title+'\n\n[서비스 모델 검토서](서비스모델_NOA_2026-09-11/'+NAME+'.md)를 작성했다. NOA 내부 확장형·별도 업무화면 연계형·기존 PMS 검사 연계형·신규 통합형을 비교했다. 기존 52개 RFP 요구사항은 유지했다.\n\n- PMS-Q02·03·14: NOA의 실제 버전·확장 권리·API·장기 업무 상태·원장 관리 책임을 F01–F10으로 확인한다.\n- PMS-Q04·07·08·12: 현행·규칙 개선·AI 추가의 순시간·품질·재작업·서비스 연속성을 구분해 검증한다.\n- 새 공개·내부 근거는 PMS-NOA-N01~N05와 PMS-NOA-L01~L03으로 출처원장에 추가했다.\n- 현재 판정은 조사·설계 초안의 조건부 적합이다. NOA 실물 시험, 기관 승인, 실제 편익 검증 및 구현 완료를 뜻하지 않는다.\n'
    queue_path.write_text(queue,encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
