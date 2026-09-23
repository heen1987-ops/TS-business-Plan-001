from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parent.parent
FONT='맑은 고딕'; BLACK=RGBColor(0,0,0); OUTS={}; bookmark_seq=0
LINK_MAP={'요구사항_분석서.md':'Public_AI_PMS_요구사항_분석서_v0.1.docx','요구사항_명세서.md':'Public_AI_PMS_요구사항_명세서_v0.1.docx'}

def font(r,size=None,bold=None):
    r.font.name=FONT
    pr=r._element.get_or_add_rPr();f=pr.find(qn('w:rFonts'))
    if f is None:f=OxmlElement('w:rFonts');pr.insert(0,f)
    for k in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+k),FONT)
    if size:r.font.size=Pt(size)
    if bold is not None:r.bold=bold
    r.font.color.rgb=BLACK

def bookmark(p,name):
    global bookmark_seq
    bookmark_seq+=1
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bookmark_seq));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(bookmark_seq))
    p._p.insert(0,a);p._p.append(b)

def link(p,label,target):
    h=OxmlElement('w:hyperlink')
    if target.startswith('#'):h.set(qn('w:anchor'),target[1:].replace('-','_'))
    else:
        raw,sep,anchor=target.partition('#')
        if raw in LINK_MAP:raw=LINK_MAP[raw]
        url=Path(raw).as_uri() if raw.startswith('C:/') else (ROOT/raw).resolve().as_uri()
        if anchor:url+='#'+anchor.replace('-','_')
        h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');ff=OxmlElement('w:rFonts')
    for k in ['ascii','hAnsi','eastAsia']:ff.set(qn('w:'+k),FONT)
    pr.append(ff);color=OxmlElement('w:color');color.set(qn('w:val'),'1F4E79');pr.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    rr.append(pr);t=OxmlElement('w:t');t.text=label;rr.append(t);h.append(rr);p._p.append(h)

def inline(p,text,size=None,bold=False):
    text=text.replace('<br>','\n').replace('\\|','|')
    pat=re.compile(r'\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|'+chr(96)+'([^'+chr(96)+']+)'+chr(96))
    pos=0
    for mt in pat.finditer(text):
        if mt.start()>pos:font(p.add_run(text[pos:mt.start()]),size,bold)
        if mt.group(1):link(p,mt.group(1),mt.group(2))
        else:font(p.add_run(mt.group(3) or mt.group(4)),size,True if mt.group(3) else bold)
        pos=mt.end()
    if pos<len(text):font(p.add_run(text[pos:]),size,bold)

def clean_heading(s):
    return re.sub(r'\s+',' ',re.sub(r'[-·/.:()—+→↔]',' ',s)).strip()

def setup(title):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.left_margin=sec.right_margin=Inches(.72)
    sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65)
    sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
    for s in doc.styles:
        if s.type==1:
            s.font.name=FONT;s.font.color.rgb=BLACK
            pr=s.element.get_or_add_rPr();f=pr.find(qn('w:rFonts'))
            if f is None:f=OxmlElement('w:rFonts');pr.insert(0,f)
            for k in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+k),FONT)
    n=doc.styles['Normal'];n.font.size=Pt(11)
    n.paragraph_format.line_spacing=1.22;n.paragraph_format.space_after=Pt(6)
    n.paragraph_format.widow_control=True
    for name,size in [('Title',26),('Subtitle',12),('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
        s=doc.styles[name];s.font.size=Pt(size);s.font.color.rgb=BLACK;s.font.bold=name!='Subtitle'
        s.paragraph_format.space_before=Pt(13 if name!='Title' else 0)
        s.paragraph_format.space_after=Pt(7);s.paragraph_format.keep_with_next=True
        pr=s.element.get_or_add_pPr()
        for b in list(pr.findall(qn('w:pBdr'))):pr.remove(b)
    for name in ['List Bullet','List Number']:
        doc.styles[name].paragraph_format.space_after=Pt(5)
        doc.styles[name].paragraph_format.line_spacing=1.22
    for nm in ['TOC 1','TOC 2']:
        try:s=doc.styles[nm]
        except KeyError:s=doc.styles.add_style(nm,1)
        s.font.name=FONT;s.font.size=Pt(11);s.font.color.rgb=BLACK
        s.paragraph_format.space_after=Pt(7)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run(title+'   |   '),8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='1';rr.append(tt);field.append(rr);footer._p.append(field)
    up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');doc.settings.element.append(up)
    doc.core_properties.title=title
    doc.core_properties.subject='공공 정보화사업 관리 및 검증 서비스 요구사항'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.keywords='Public AI-PMS, 요구사항, 설계'
    return doc

def field_toc(doc):
    style=doc.styles.add_style('문서 목차 제목',1)
    style.font.name=FONT;style.font.size=Pt(16);style.font.bold=True;style.font.color.rgb=BLACK
    style.paragraph_format.space_before=Pt(14);style.paragraph_format.space_after=Pt(8)
    p=doc.add_paragraph('목차',style='문서 목차 제목')
    out=OxmlElement('w:outlineLvl');out.set(qn('w:val'),'9');p._p.get_or_add_pPr().append(out)
    pp=doc.add_paragraph()
    for typ in ['begin','instruction','separate','end']:
        a=OxmlElement('w:r')
        if typ=='instruction':
            f=OxmlElement('w:instrText');f.set(qn('xml:space'),'preserve');f.text=' TOC \\o "1-1" \\h \\z \\u '
        else:
            f=OxmlElement('w:fldChar');f.set(qn('w:fldCharType'),typ)
        a.append(f);pp._p.append(a)
        if typ=='separate':font(pp.add_run('목차'),11)

def widths(headers):
    n=len(headers);w=508.32
    if n==2:return [99,w-99]
    if headers[0].startswith('요구 ID'):return [120,235,55,52,w-462]
    if headers[0].startswith('역할 ID'):return [45,105,179,w-329]
    if headers[0].startswith('자료 ID'):return [101,246,w-347]
    if n==3:return [103,211,w-314]
    if n==4:return [75,151,151,w-377]
    if n==5:return [66,115,120,115,w-416]
    if n==6:return [65,93,93,93,92,w-436]
    return [w/n]*n

def table(doc,rs):
    n=len(rs[0]);ww=widths(rs[0])
    t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for k in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for c,v in zip(t.columns,ww):c.width=Pt(v)
    for i,row in enumerate(rs):
        cells=t.add_row().cells;tr=t.rows[-1]._tr.get_or_add_trPr()
        tr.append(OxmlElement('w:cantSplit'))
        if i==0:
            rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');tr.append(rep)
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Pt(ww[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for tag,val in [('top',65),('bottom',65),('left',85),('right',85)]:
                z=OxmlElement('w:'+tag);z.set(qn('w:w'),str(val));z.set(qn('w:type'),'dxa');margin.append(z)
            cp.append(margin)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EDF4');cp.append(shade)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.10
            p.paragraph_format.keep_with_next=i==0
            inline(p,value,10 if n<=3 else 9.5,i==0)
    gap=doc.add_paragraph();gap.paragraph_format.space_after=Pt(3)
    gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.line_spacing=1
    font(gap.add_run(''),2)
    return t

def parse_table(lines,i):
    rs=[]
    while i<len(lines) and lines[i].strip().startswith('|'):
        cols=[x.strip().replace('\\|','|') for x in re.split(r'(?<!\\)\|',lines[i].strip())[1:-1]]
        if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in cols):rs.append(cols)
        i+=1
    return rs,i

def convert(name,outname):
    source=ROOT/name;raw=source.read_text(encoding='utf-8');lines=raw.splitlines()
    title='Public AI PMS '+('요구사항 분석서' if '분석서' in name else '요구사항 명세서')
    doc=setup(title);intro_end=next(i for i,l in enumerate(lines) if l.startswith('## '))
    doc.add_paragraph(title,style='Title')
    doc.add_paragraph('공공 정보화사업 이행관리와 검증 지원 서비스',style='Subtitle')
    for l in lines[1:intro_end]:
        if l.strip():inline(doc.add_paragraph(),l)
    doc.add_paragraph('기관 검토와 서비스 범위 협의를 위한 요구사항 문서')
    field_toc(doc);doc.add_page_break()
    i=intro_end;cards=0
    while i<len(lines):
        line=lines[i].strip()
        if not line or line=='---' or re.match(r'<a id="([^"]+)"></a>',line):i+=1;continue
        if line.startswith('|'):
            rs,i=parse_table(lines,i);table(doc,rs);continue
        h=re.match(r'^(#{2,4})\s+(.*)',line)
        if h:
            level=len(h.group(1))-1;text=h.group(2)
            card=re.match(r'(PMS-(?:SFR|NFR|DAR|IFR|OPR|BMR)-\d{3}) · (.*)',text)
            if card:
                p=doc.add_paragraph(clean_heading(card.group(2)),style='Heading 2')
                p.paragraph_format.page_break_before=cards>0
                bookmark(p,card.group(1).lower().replace('-','_'))
                p2=doc.add_paragraph();p2.paragraph_format.keep_with_next=True
                inline(p2,'**'+card.group(1)+'**',11);cards+=1
            else:
                p=doc.add_paragraph(clean_heading(text),style='Heading '+str(min(level,3)))
                if level==1 and '명세서' in name and text.startswith(('9.','10.')):p.paragraph_format.page_break_before=True
            i+=1;continue
        if line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\.\s',line):
            p=doc.add_paragraph();inline(p,line)
            p.paragraph_format.left_indent=Pt(14);p.paragraph_format.first_line_indent=Pt(-14)
        else:
            p=doc.add_paragraph();inline(p,line)
            if line=='**수용기준 — 미실행**':p.paragraph_format.keep_with_next=True
        i+=1
    out=ROOT/outname;doc.save(out)
    OUTS[outname]={'source':name,'sha256_source':hashlib.sha256(source.read_bytes()).hexdigest(),'cards':cards,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)}
for source,out in LINK_MAP.items():convert(source,out)
(ROOT/'_word_build'/'build_manifest.json').write_text(json.dumps(OUTS,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(OUTS,ensure_ascii=False,indent=2))

